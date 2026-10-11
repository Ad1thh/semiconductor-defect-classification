import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TEMPLATE_PATH = r"C:\Users\noopa\Downloads\Mini Project PPT format (1).pptx"
OUT_PATHS = [
    r"c:\Users\noopa\Downloads\semiconductor-defect-classification\Thank_You_Slide.pptx",
    r"C:\Users\noopa\Downloads\Thank_You_Slide.pptx"
]

prs = pptx.Presentation(TEMPLATE_PATH)
DARK_NAVY = RGBColor(15, 23, 42)
TEXT_DARK = RGBColor(30, 41, 59)
BLUE_ACCENT = RGBColor(2, 132, 199)

# Delete slides 0 to 13 (leaving only slide 14, the 15th slide - Thank You)
total_slides = len(prs.slides)
for i in range(total_slides - 2, -1, -1):
    rId = prs.slides._sldIdLst[i].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[i]

print(f"Remaining slides: {len(prs.slides)}")
thank_you_slide = prs.slides[0]

# Format Thank You slide
for shape in thank_you_slide.shapes:
    if shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if "THANK YOU" in t:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = "THANK YOU"
            run.font.name = "Times New Roman"
            run.font.size = Pt(56)
            run.font.bold = True
            run.font.color.rgb = DARK_NAVY
        elif "Questions" in t:
            tf = shape.text_frame
            tf.clear()
            lines = [
                ("Questions & Discussion", 28, True, BLUE_ACCENT),
                ("", 12, False, None),
                ("Automated Semiconductor Wafer Defect Inspection System", 22, True, DARK_NAVY),
                ("B.Tech Mini Project Final Internal Review • Saintgits College of Engineering (Autonomous)", 18, False, TEXT_DARK),
                ("Team: Adithyan B (MGP23EVLSI001) • Anna Boban (MGP23EVLSI002) • Noopa Maria Rojan (MGP23EVLSI009)", 16, False, TEXT_DARK),
                ("Supervisor: Er. Justin Mathew, Assistant Professor, Dept. of ECE", 16, False, TEXT_DARK)
            ]
            for idx, (txt, sz, bld, col) in enumerate(lines):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                if txt:
                    r = p.add_run()
                    r.text = txt
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(sz)
                    r.font.bold = bld
                    if col:
                        r.font.color.rgb = col
        elif t.isdigit():
            # Clear old slide number 15
            shape.text_frame.clear()

for out_path in OUT_PATHS:
    prs.save(out_path)
    print(f"Saved: {out_path} ({os.path.getsize(out_path)} bytes)")
